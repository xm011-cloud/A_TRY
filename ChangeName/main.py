#!/usr/bin/env python3
"""
批量文件重命名工具 - 图形界面版
支持拖拽文件/文件夹，多种重命名规则
窗口布局：可拖拽分隔条自由调整各区域高度
预览功能：文本、图片、Word/PDF
"""

import os
import re
import subprocess
import sys
from pathlib import Path
from typing import List, Tuple
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# 可选依赖
try:
    from PIL import Image, ImageTk
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


# PDF依赖
try:
    import PyPDF2
    HAS_PYPDF2 = True
except ImportError:
    HAS_PYPDF2 = False
try:
    from pdf2image import convert_from_path
    HAS_PDF2IMG = True
except ImportError:
    HAS_PDF2IMG = False


# 拖拽支持优化，防止导入失败导致崩溃
HAS_DND = False
try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    HAS_DND = True
except ImportError:
    class DummyDnD:
        def __init__(self, *args, **kwargs):
            self._tk = tk.Tk(*args, **kwargs)
        def __getattr__(self, item):
            return getattr(self._tk, item)
    DND_FILES = None
    TkinterDnD = DummyDnD



class FileRenamerGUI:
    def __init__(self):
        try:
            self.root = TkinterDnD.Tk() if HAS_DND else tk.Tk()
        except Exception as e:
            messagebox.showerror("错误", f"无法初始化主窗口: {e}\n请检查Python环境和tkinter安装。")
            sys.exit(1)
        self.root.title("批量文件重命名工具")
        self.root.geometry("1000x750")
        self.root.minsize(800, 600)

        self.files: List[Path] = []
        self.preview_data: List[Tuple[Path, Path]] = []

        self.create_widgets()
        try:
            self.setup_drag_drop()
        except Exception as e:
            # 拖拽失败不影响主功能
            pass
        self.bind_events()
        self.update_preview()

        self.root.mainloop()

    # ------------------ 界面构建（使用 PanedWindow + grid）------------------
    def create_widgets(self):
        # 主容器：使用 grid 布局，确保底部按钮始终可见
        main_container = ttk.Frame(self.root)
        main_container.pack(fill=tk.BOTH, expand=True)

        # 配置 grid 行权重
        main_container.rowconfigure(0, weight=1)   # PanedWindow 区域
        main_container.rowconfigure(1, weight=0)   # 按钮栏
        main_container.rowconfigure(2, weight=0)   # 状态栏
        main_container.columnconfigure(0, weight=1)

        # ---------- 1. 可拖拽分隔条区域 ----------
        main_pane = ttk.PanedWindow(main_container, orient=tk.VERTICAL)
        main_pane.grid(row=0, column=0, sticky="nsew", pady=(0, 5))

        # 文件列表框架
        list_frame = ttk.LabelFrame(main_pane, text="文件列表 (支持拖拽，双击预览)", padding="5")
        main_pane.add(list_frame, weight=3)

        # 文件列表内部布局
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(1, weight=1)

        btn_frame = ttk.Frame(list_frame)
        btn_frame.grid(row=0, column=0, sticky="ew", pady=(0, 5))
        btn_frame.columnconfigure(7, weight=1)

        ttk.Button(btn_frame, text="添加文件", command=self.add_files).grid(row=0, column=0, padx=2)
        ttk.Button(btn_frame, text="添加文件夹", command=self.add_folder).grid(row=0, column=1, padx=2)
        ttk.Button(btn_frame, text="清空全部", command=self.clear_files).grid(row=0, column=2, padx=2)
        ttk.Button(btn_frame, text="移除选中", command=self.remove_selected).grid(row=0, column=3, padx=2)

        self.recursive_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(btn_frame, text="递归子目录", variable=self.recursive_var).grid(row=0, column=4, padx=(10, 0))

        list_container = ttk.Frame(list_frame)
        list_container.grid(row=1, column=0, sticky="nsew")
        list_container.columnconfigure(0, weight=1)
        list_container.rowconfigure(0, weight=1)

        self.file_listbox = tk.Listbox(list_container, selectmode=tk.EXTENDED)
        scrollbar = ttk.Scrollbar(list_container, orient=tk.VERTICAL, command=self.file_listbox.yview)
        self.file_listbox.configure(yscrollcommand=scrollbar.set)
        self.file_listbox.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.file_listbox.bind("<Double-Button-1>", self.on_file_double_click)

        # ---------- 规则区域 ----------
        rule_frame = ttk.LabelFrame(main_pane, text="重命名规则", padding="10")
        main_pane.add(rule_frame, weight=2)

        # 规则区域网格布局（宽度可调）
        for i in range(6):
            rule_frame.columnconfigure(i, weight=1 if i % 2 == 1 else 0)

        row = 0
        # 替换
        ttk.Label(rule_frame, text="替换:").grid(row=row, column=0, sticky="w", padx=5, pady=2)
        self.replace_old = tk.StringVar()
        self.replace_new = tk.StringVar()
        ttk.Entry(rule_frame, textvariable=self.replace_old).grid(row=row, column=1, sticky="ew", padx=5, pady=2)
        ttk.Label(rule_frame, text="→").grid(row=row, column=2, padx=5)
        ttk.Entry(rule_frame, textvariable=self.replace_new).grid(row=row, column=3, sticky="ew", padx=5, pady=2)
        row += 1

        # 前缀/后缀
        ttk.Label(rule_frame, text="前缀:").grid(row=row, column=0, sticky="w", padx=5, pady=2)
        self.prefix = tk.StringVar()
        ttk.Entry(rule_frame, textvariable=self.prefix).grid(row=row, column=1, sticky="ew", padx=5, pady=2)
        ttk.Label(rule_frame, text="后缀:").grid(row=row, column=2, sticky="w", padx=5, pady=2)
        self.suffix = tk.StringVar()
        ttk.Entry(rule_frame, textvariable=self.suffix).grid(row=row, column=3, sticky="ew", padx=5, pady=2)
        row += 1

        # 正则替换
        ttk.Label(rule_frame, text="正则替换:").grid(row=row, column=0, sticky="w", padx=5, pady=2)
        self.regex_pattern = tk.StringVar()
        self.regex_repl = tk.StringVar()
        ttk.Entry(rule_frame, textvariable=self.regex_pattern).grid(row=row, column=1, sticky="ew", padx=5, pady=2)
        ttk.Label(rule_frame, text="→").grid(row=row, column=2, padx=5)
        ttk.Entry(rule_frame, textvariable=self.regex_repl).grid(row=row, column=3, sticky="ew", padx=5, pady=2)
        self.ignore_case = tk.BooleanVar(value=False)
        ttk.Checkbutton(rule_frame, text="忽略大小写", variable=self.ignore_case).grid(row=row, column=4, columnspan=2, sticky="w", padx=5)
        row += 1

        # 序号模式
        ttk.Label(rule_frame, text="序号模式:").grid(row=row, column=0, sticky="w", padx=5, pady=2)
        self.number_pattern = tk.StringVar()
        ttk.Entry(rule_frame, textvariable=self.number_pattern).grid(row=row, column=1, columnspan=3, sticky="ew", padx=5, pady=2)
        ttk.Label(rule_frame, text="(例如: photo_{:03d} → photo_001.jpg)").grid(row=row, column=4, columnspan=2, sticky="w", padx=5)
        row += 1

        # 扩展名过滤
        ttk.Label(rule_frame, text="仅处理扩展名:").grid(row=row, column=0, sticky="w", padx=5, pady=2)
        self.ext_filter = tk.StringVar()
        ttk.Entry(rule_frame, textvariable=self.ext_filter).grid(row=row, column=1, columnspan=3, sticky="ew", padx=5, pady=2)
        ttk.Label(rule_frame, text="(多个用逗号分隔，如 .jpg,.png)").grid(row=row, column=4, columnspan=2, sticky="w", padx=5)
        row += 1

        # 文件名筛选
        ttk.Label(rule_frame, text="筛选文件名:").grid(row=row, column=0, sticky="w", padx=5, pady=2)
        self.name_filter = tk.StringVar()
        ttk.Entry(rule_frame, textvariable=self.name_filter).grid(row=row, column=1, columnspan=3, sticky="ew", padx=5, pady=2)
        ttk.Label(rule_frame, text="(正则匹配，不填则不过滤)").grid(row=row, column=4, columnspan=2, sticky="w", padx=5)

        # ---------- 重命名预览区域 ----------
        preview_frame = ttk.LabelFrame(main_pane, text="预览 (将重命名的文件)", padding="5")
        main_pane.add(preview_frame, weight=2)

        preview_frame.columnconfigure(0, weight=1)
        preview_frame.rowconfigure(0, weight=1)

        self.preview_text = tk.Text(preview_frame, wrap=tk.NONE, font=("Consolas", 9))
        vsb = ttk.Scrollbar(preview_frame, orient=tk.VERTICAL, command=self.preview_text.yview)
        hsb = ttk.Scrollbar(preview_frame, orient=tk.HORIZONTAL, command=self.preview_text.xview)
        self.preview_text.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self.preview_text.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")

        # ---------- 2. 底部按钮栏 ----------
        action_frame = ttk.Frame(main_container)
        action_frame.grid(row=1, column=0, sticky="ew", padx=10, pady=(0, 5))
        action_frame.columnconfigure(2, weight=1)

        ttk.Button(action_frame, text="刷新预览", command=self.update_preview).grid(row=0, column=0, padx=5)
        ttk.Button(action_frame, text="执行重命名", command=self.execute_rename).grid(row=0, column=1, padx=5)

        # ---------- 3. 状态栏 ----------
        self.status_var = tk.StringVar(value="就绪")
        status_bar = ttk.Label(main_container, textvariable=self.status_var, relief=tk.SUNKEN, anchor=tk.W)
        status_bar.grid(row=2, column=0, sticky="ew")

    def bind_events(self):
        vars_to_trace = [
            self.replace_old, self.replace_new, self.prefix, self.suffix,
            self.regex_pattern, self.regex_repl, self.number_pattern,
            self.ext_filter, self.name_filter
        ]
        for var in vars_to_trace:
            var.trace_add('write', lambda *_: self.update_preview())
        self.ignore_case.trace_add('write', lambda *_: self.update_preview())

    def setup_drag_drop(self):
        if HAS_DND and DND_FILES is not None:
            try:
                self.file_listbox.drop_target_register(DND_FILES)
                self.file_listbox.dnd_bind('<<Drop>>', self.on_drop)
                self.root.drop_target_register(DND_FILES)
                self.root.dnd_bind('<<Drop>>', self.on_drop)
            except Exception:
                pass

    # ------------------ 文件管理 ------------------
    def on_drop(self, event):
        raw_data = event.data
        paths = self.parse_drop_data(raw_data)
        for p in paths:
            self.add_path(p)
        self.update_preview()

    def parse_drop_data(self, data: str) -> List[str]:
        import shlex
        data = data.strip()
        if data.startswith('{') and data.endswith('}'):
            data = data[1:-1]
        try:
            parts = shlex.split(data)
        except:
            parts = data.split()
        return [os.path.normpath(p) for p in parts]

    def add_files(self):
        paths = filedialog.askopenfilenames(title="选择文件")
        for p in paths:
            self.add_path(p)
        self.update_preview()

    def add_folder(self):
        folder = filedialog.askdirectory(title="选择文件夹")
        if not folder:
            return
        self.add_path(folder, is_folder=True)
        self.update_preview()

    def add_path(self, path: str, is_folder: bool = False):
        try:
            p = Path(path)
            if not p.exists():
                return
            if is_folder or p.is_dir():
                recursive = self.recursive_var.get()
                files = list(p.rglob("*")) if recursive else list(p.glob("*"))
                files = [f for f in files if f.is_file()]
                self.files.extend(files)
            else:
                self.files.append(p)
            # 去重
            self.files = list({str(f): f for f in self.files}.values())
        except Exception as e:
            messagebox.showerror("添加文件失败", f"路径: {path}\n错误: {e}")

    def clear_files(self):
        self.files.clear()
        self.update_preview()

    def remove_selected(self):
        selected = self.file_listbox.curselection()
        for idx in reversed(selected):
            if idx < len(self.files):
                del self.files[idx]
        self.update_preview()

    # ------------------ 文件内容预览（增强）------------------
    def on_file_double_click(self, event):
        selection = self.file_listbox.curselection()
        if not selection:
            return
        idx = selection[0]
        if idx < len(self.files):
            self.show_file_content(self.files[idx])

    def show_file_content(self, file_path: Path):
        ext = file_path.suffix.lower()
        try:
            stat = file_path.stat()
            size = stat.st_size
            mtime = stat.st_mtime
        except Exception as e:
            messagebox.showerror("错误", f"无法读取文件信息: {e}")
            return

        # 图片预览
        if ext in {'.jpg', '.jpeg', '.png', '.gif', '.bmp', '.ico'} and HAS_PIL:
            self.preview_image(file_path)
            return

        # Word 预览
        if ext == '.docx' and HAS_DOCX:
            self.preview_docx(file_path)
            return

        # PDF 预览（优先图片）
        if ext == '.pdf' and (HAS_PDF2IMG or HAS_PYPDF2):
            if HAS_PDF2IMG:
                try:
                    self.preview_pdf_image(file_path)
                    return
                except Exception as e:
                    # 回退文本
                    pass
            if HAS_PYPDF2:
                self.preview_pdf(file_path)
                return

        # 文本或二进制
        is_text, content = self.read_file_content(file_path, preview_limit=64*1024)
        if is_text:
            self.preview_text_file(file_path, content, size, mtime)
        else:
            self.preview_binary_file(file_path, size, mtime)

    def preview_pdf_image(self, file_path: Path):
        # 用pdf2image渲染第一页为图片
        try:
            images = convert_from_path(str(file_path), first_page=1, last_page=1)
            if not images:
                raise RuntimeError("未能渲染PDF页面")
            img = images[0]
            win = tk.Toplevel(self.root)
            win.title(f"PDF图片预览 - {file_path.name}")
            win.geometry("800x600")
            win.minsize(400, 300)
            info = f"PDF第一页图片预览 | 尺寸: {img.width} x {img.height}"
            ttk.Label(win, text=info, padding=5).pack(anchor=tk.W)
            max_w, max_h = 700, 500
            img.thumbnail((max_w, max_h))
            photo = ImageTk.PhotoImage(img)
            label = ttk.Label(win, image=photo)
            label.image = photo
            label.pack(expand=True, fill=tk.BOTH, padx=10, pady=10)
        except Exception as e:
            messagebox.showerror("PDF图片预览失败", f"无法渲染PDF为图片: {e}")

    def preview_image(self, file_path: Path):
        try:
            img = Image.open(file_path)
            win = tk.Toplevel(self.root)
            win.title(f"图片预览 - {file_path.name}")
            win.geometry("800x600")
            win.minsize(400, 300)

            info = f"尺寸: {img.width} x {img.height} | 格式: {img.format}"
            ttk.Label(win, text=info, padding=5).pack(anchor=tk.W)

            max_w, max_h = 700, 500
            img.thumbnail((max_w, max_h), Image.Resampling.LANCZOS)
            photo = ImageTk.PhotoImage(img)
            label = ttk.Label(win, image=photo)
            label.image = photo
            label.pack(expand=True, fill=tk.BOTH, padx=10, pady=10)
        except Exception as e:
            messagebox.showerror("预览错误", f"无法打开图片: {e}")

    def preview_docx(self, file_path: Path):
        try:
            doc = docx.Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
            if not text.strip():
                text = "文档无文本内容（可能全是图片或表格）"
        except Exception as e:
            messagebox.showerror("读取失败", f"无法解析 Word 文档: {e}")
            return
        win = self._create_text_preview_window(f"Word文档预览 - {file_path.name}", text)
        ttk.Label(win, text="文档文本内容（仅文本，不含格式）", foreground="blue").pack(anchor=tk.W)

    def preview_pdf(self, file_path: Path):
        try:
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                text = ""
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                if not text.strip():
                    text = "PDF 中未提取到文本（可能是扫描图片）"
        except Exception as e:
            messagebox.showerror("读取失败", f"无法解析 PDF: {e}")
            return
        win = self._create_text_preview_window(f"PDF预览 - {file_path.name}", text)
        ttk.Label(win, text="提取的文本内容（可能不完整）", foreground="blue").pack(anchor=tk.W)

    def preview_text_file(self, file_path: Path, content: str, size: int, mtime: float):
        win = self._create_text_preview_window(f"文件预览 - {file_path.name}", content)
        info_frame = ttk.Frame(win)
        info_frame.pack(fill=tk.X, pady=(0,5))
        size_str = self.format_size(size)
        mtime_str = self.format_time(mtime)
        info_text = f"路径: {file_path}\n大小: {size_str} | 修改时间: {mtime_str}"
        if size > 64*1024:
            info_text += "\n⚠️ 文件较大，仅显示前 64KB 内容（完整文件请用外部编辑器打开）"
        ttk.Label(info_frame, text=info_text, wraplength=780, justify=tk.LEFT).pack(anchor=tk.W)

    def preview_binary_file(self, file_path: Path, size: int, mtime: float):
        win = tk.Toplevel(self.root)
        win.title(f"文件预览 - {file_path.name}")
        win.geometry("500x300")
        win.minsize(400, 250)

        info_frame = ttk.Frame(win, padding=10)
        info_frame.pack(fill=tk.X)
        info_text = f"文件类型: {file_path.suffix or '未知'}\n大小: {self.format_size(size)}\n修改时间: {self.format_time(mtime)}\n无法内置预览（二进制文件或未安装相应库）。"
        ttk.Label(info_frame, text=info_text, wraplength=460).pack(anchor=tk.W, pady=5)

        def open_external():
            if messagebox.askyesno("安全确认", f"即将使用系统默认程序打开文件：\n{file_path.name}\n\n请确保文件来源可信。是否继续？", icon='warning'):
                self.open_with_default_app(file_path)
                win.destroy()

        ttk.Button(win, text="用默认程序打开", command=open_external).pack(pady=10)

        ext = file_path.suffix.lower()
        if ext == '.docx' and not HAS_DOCX:
            ttk.Label(win, text="提示：安装 python-docx 可预览 Word 文档\n命令: pip install python-docx", foreground="gray").pack()
        elif ext == '.pdf' and not HAS_PYPDF2:
            ttk.Label(win, text="提示：安装 PyPDF2 可预览 PDF 文档\n命令: pip install PyPDF2", foreground="gray").pack()
        elif ext in {'.jpg', '.png'} and not HAS_PIL:
            ttk.Label(win, text="提示：安装 Pillow 可预览图片\n命令: pip install Pillow", foreground="gray").pack()

    def _create_text_preview_window(self, title: str, content: str) -> tk.Toplevel:
        win = tk.Toplevel(self.root)
        win.title(title)
        win.geometry("800x600")
        win.minsize(400, 300)

        text_frame = ttk.Frame(win)
        text_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        text_widget = tk.Text(text_frame, wrap=tk.NONE, font=("Consolas", 10))
        vsb = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=text_widget.yview)
        hsb = ttk.Scrollbar(text_frame, orient=tk.HORIZONTAL, command=text_widget.xview)
        text_widget.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        text_widget.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        text_frame.columnconfigure(0, weight=1)
        text_frame.rowconfigure(0, weight=1)

        text_widget.insert(tk.END, content)
        text_widget.configure(state=tk.DISABLED)
        return win

    def read_file_content(self, file_path: Path, preview_limit=64*1024) -> Tuple[bool, str]:
        file_size = file_path.stat().st_size
        is_large = file_size > preview_limit

        text_exts = {
            '.txt', '.py', '.js', '.html', '.css', '.json', '.xml', '.csv', '.md',
            '.ini', '.cfg', '.conf', '.log', '.bat', '.sh', '.c', '.cpp', '.h', '.java',
            '.go', '.rs', '.swift', '.kt', '.rb', '.pl', '.php', '.sql', '.yaml', '.yml'
        }

        def try_read(limit_bytes=None):
            encodings = ['utf-8', 'gbk', 'latin-1']
            for enc in encodings:
                try:
                    if limit_bytes:
                        with open(file_path, 'r', encoding=enc) as f:
                            content = f.read(limit_bytes)
                        if is_large and limit_bytes:
                            try:
                                encoded_len = len(content.encode(enc))
                            except:
                                encoded_len = limit_bytes
                            if encoded_len < file_size:
                                remaining = file_size - encoded_len
                                content += f"\n\n... (文件过大，仅显示前 {encoded_len//1024} KB，剩余 {remaining//1024} KB)"
                        return True, content
                    else:
                        with open(file_path, 'r', encoding=enc) as f:
                            content = f.read()
                        return True, content
                except (UnicodeDecodeError, OSError):
                    continue
            return False, "无法解码文件内容，可能是二进制文件或未知编码。"

        if file_path.suffix.lower() in text_exts:
            return try_read(limit_bytes=preview_limit if is_large else None)

        try:
            with open(file_path, 'rb') as f:
                sample = f.read(512)
            if b'\x00' in sample:
                return False, "文件为二进制文件，无法显示文本内容。"
            return try_read(limit_bytes=preview_limit if is_large else None)
        except Exception as e:
            return False, f"读取文件时出错: {e}"

    def open_with_default_app(self, path: Path):
        try:
            if sys.platform == 'win32':
                os.startfile(path)
            elif sys.platform == 'darwin':
                subprocess.Popen(['open', path])
            else:
                subprocess.Popen(['xdg-open', path])
        except Exception as e:
            messagebox.showerror("打开失败", f"无法用默认程序打开文件: {e}")

    def format_size(self, size: int) -> str:
        for unit in ['B', 'KB', 'MB', 'GB']:
            if size < 1024.0:
                return f"{size:.1f} {unit}"
            size /= 1024.0
        return f"{size:.1f} TB"

    def format_time(self, timestamp: float) -> str:
        import datetime
        dt = datetime.datetime.fromtimestamp(timestamp)
        return dt.strftime("%Y-%m-%d %H:%M:%S")

    # ------------------ 重命名核心 ------------------
    def update_preview(self):
        self.file_listbox.delete(0, tk.END)
        for f in self.files:
            self.file_listbox.insert(tk.END, str(f))

        self.preview_data.clear()
        if not self.files:
            self.preview_text.delete(1.0, tk.END)
            self.preview_text.insert(tk.END, "无文件")
            self.status_var.set("文件数: 0")
            return

        ext_filter_str = self.ext_filter.get().strip()
        ext_list = []
        if ext_filter_str:
            ext_list = [e.strip().lower() for e in ext_filter_str.split(',')]
            ext_list = [e if e.startswith('.') else f'.{e}' for e in ext_list]

        name_pattern = self.name_filter.get().strip()
        name_regex = None
        if name_pattern:
            flags = re.IGNORECASE if self.ignore_case.get() else 0
            try:
                name_regex = re.compile(name_pattern, flags)
            except re.error:
                self.preview_text.delete(1.0, tk.END)
                self.preview_text.insert(tk.END, "错误：筛选正则无效")
                return

        rename_func = self.build_rename_function()

        filtered = []
        for src in self.files:
            if ext_list and src.suffix.lower() not in ext_list:
                continue
            if name_regex and not name_regex.search(src.name):
                continue
            filtered.append(src)

        filtered.sort(key=lambda p: p.name)

        lines = []
        for idx, src in enumerate(filtered, 1):
            try:
                new_name = rename_func(src, idx)
                if not new_name:
                    new_name = src.name
                dst = src.parent / new_name
                self.preview_data.append((src, dst))
                lines.append(f"{src.name}  →  {dst.name}")
            except Exception as e:
                lines.append(f"{src.name}  →  错误: {e}")

        self.preview_text.delete(1.0, tk.END)
        if lines:
            self.preview_text.insert(tk.END, "\n".join(lines))
        else:
            self.preview_text.insert(tk.END, "没有符合过滤条件的文件")

        self.status_var.set(f"文件总数: {len(self.files)} | 符合条件: {len(filtered)}")

    def build_rename_function(self):
        rules = []

        old = self.replace_old.get()
        new = self.replace_new.get()
        if old:
            rules.append(lambda name: name.replace(old, new))

        prefix = self.prefix.get()
        suffix = self.suffix.get()

        if prefix:
            rules.append(lambda name: prefix + name)
        if suffix:
            def add_suffix(name):
                stem = Path(name).stem
                ext = Path(name).suffix
                return f"{stem}{suffix}{ext}"
            rules.append(add_suffix)

        regex_pat = self.regex_pattern.get()
        regex_repl = self.regex_repl.get()
        if regex_pat:
            flags = re.IGNORECASE if self.ignore_case.get() else 0
            try:
                regex = re.compile(regex_pat, flags)
                rules.append(lambda name: regex.sub(regex_repl, name))
            except re.error:
                pass

        number_pat = self.number_pattern.get().strip()
        if number_pat:
            def number_rename(src: Path, idx: int) -> str:
                try:
                    return number_pat.format(idx) + src.suffix
                except Exception:
                    return src.name
            return number_rename

        def combined_rename(src: Path, idx: int) -> str:
            name = src.name
            for rule in rules:
                try:
                    name = rule(name)
                except Exception:
                    continue
            return name
        return combined_rename

    def execute_rename(self):
        if not self.preview_data:
            messagebox.showinfo("提示", "没有要重命名的文件")
            return
        if not messagebox.askyesno("确认重命名", f"即将重命名 {len(self.preview_data)} 个文件。\n是否继续？", icon='warning'):
            return

        conflicts = [(src, dst) for src, dst in self.preview_data if dst.exists() and dst != src]
        if conflicts:
            msg = "以下目标文件已存在，将被覆盖：\n" + "\n".join(f"{src.name} -> {dst.name}" for src, dst in conflicts[:10])
            if len(conflicts) > 10:
                msg += f"\n...等{len(conflicts)}个"
            msg += "\n\n是否继续？"
            if not messagebox.askyesno("冲突警告", msg, icon='warning'):
                return

        success = 0
        errors = []
        for src, dst in self.preview_data:
            try:
                if dst.exists() and dst != src:
                    dst.unlink()
                src.rename(dst)
                success += 1
            except Exception as e:
                errors.append(f"{src.name} -> {dst.name}: {e}")

        new_files = [dst for _, dst in self.preview_data if dst.exists()]
        self.files = new_files
        self.update_preview()

        msg = f"完成！成功 {success} 个，失败 {len(errors)} 个。"
        if errors:
            msg += "\n错误详情：\n" + "\n".join(errors[:5])
        messagebox.showinfo("结果", msg)
        self.status_var.set(msg)


if __name__ == "__main__":
    try:
        FileRenamerGUI()
    except Exception as e:
        import traceback
        messagebox.showerror("程序异常退出", f"发生未处理异常:\n{e}\n{traceback.format_exc()}")
